"""Parse and commit admin student imports without weakening identity rules.

Preview is read-only. Commit reparses and revalidates the workbook, isolates
each selected row with a savepoint, commits all successful database work, and
only then attempts one welcome email per newly-created parent account.
Student identity is permanent; class placement is written through
``StudentClassAssignment`` for the selected school year.
"""

from __future__ import annotations

import io
import re
import secrets
from collections import Counter, defaultdict
from datetime import datetime, timezone
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from audit import create_audit_log
from auth import hash_password
from constants import normalize_class_name
from models import Class, Parent, Student, StudentParent, User
from services.email_service import send_account_created_email
from services.student_assignments import set_student_assignment


REQUIRED_COLUMNS = (
    "student_first_name",
    "student_last_name",
    "class_name",
    "parent_name",
    "parent_email",
)
OPTIONAL_COLUMNS = (
    "student_number",
    "educmaster_number",
    "parent_phone",
    "relationship",
)
ALLOWED_COLUMNS = set(REQUIRED_COLUMNS + OPTIONAL_COLUMNS)
MAX_IMPORT_ROWS = 500
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class StudentImportFileError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def _canonical_email(value: Any) -> str:
    return _text(value).lower()


def _issue(code: str, field: str | None, message: str) -> dict:
    return {"code": code, "field": field, "message": message}


def parse_student_workbook(file_bytes: bytes) -> list[dict]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise StudentImportFileError(
            "student_import_invalid_workbook",
            "The uploaded file is not a valid .xlsx workbook.",
        ) from exc

    try:
        sheet = workbook.active
        raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if raw_headers is None:
            raise StudentImportFileError("student_import_empty_workbook", "The workbook is empty.")
        headers = [_text(value) for value in raw_headers]
        nonempty_headers = [header for header in headers if header]
        duplicate_headers = sorted(header for header, count in Counter(nonempty_headers).items() if count > 1)
        if duplicate_headers:
            raise StudentImportFileError(
                "student_import_duplicate_columns",
                f"Duplicate columns: {', '.join(duplicate_headers)}",
            )
        missing = [column for column in REQUIRED_COLUMNS if column not in nonempty_headers]
        if missing:
            raise StudentImportFileError(
                "student_import_missing_columns",
                f"Missing required columns: {', '.join(missing)}",
            )
        unknown = sorted(header for header in nonempty_headers if header not in ALLOWED_COLUMNS)
        if unknown:
            raise StudentImportFileError(
                "student_import_unknown_columns",
                f"Unsupported columns: {', '.join(unknown)}",
            )

        rows: list[dict] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if any(
                not headers[index] and _text(value)
                for index, value in enumerate(values)
                if index < len(headers)
            ):
                raise StudentImportFileError(
                    "student_import_unknown_columns",
                    "Every populated column must have a supported heading.",
                )
            mapped = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            if not any(_text(value) for value in mapped.values()):
                continue
            rows.append(
                {
                    "row_number": row_number,
                    **{column: _text(mapped.get(column)) for column in ALLOWED_COLUMNS},
                }
            )
            if len(rows) > MAX_IMPORT_ROWS:
                raise StudentImportFileError(
                    "student_import_too_many_rows",
                    f"A workbook may contain at most {MAX_IMPORT_ROWS} data rows.",
                )
        if not rows:
            raise StudentImportFileError("student_import_empty_workbook", "The workbook has no student rows.")
        return rows
    finally:
        workbook.close()


def _active_user_by_canonical_email(db: Session, email: str) -> User | None:
    return db.scalar(
        select(User)
        .options(joinedload(User.parent_profile))
        .where(func.lower(User.email) == email, User.deleted_at.is_(None))
    )


def _class_maps(
    db: Session, school_year: str
) -> tuple[dict[str, list[Class]], dict[str, list[Class]]]:
    classes = db.scalars(
        select(Class).where(Class.school_year == school_year, Class.deleted_at.is_(None))
    ).all()
    french: dict[str, list[Class]] = defaultdict(list)
    english: dict[str, list[Class]] = defaultdict(list)
    for row in classes:
        french[normalize_class_name(row.name_fr)].append(row)
        if row.name_en:
            english[normalize_class_name(row.name_en)].append(row)
    return french, english


def _resolve_class(
    class_text: str,
    french: dict[str, list[Class]],
    english: dict[str, list[Class]],
) -> tuple[Class | None, dict | None]:
    key = normalize_class_name(class_text)
    french_candidates = french.get(key, [])
    if len(french_candidates) == 1:
        return french_candidates[0], None
    if len(french_candidates) > 1:
        return None, _issue(
            "student_import_class_ambiguous",
            "class_name",
            "The class label matches more than one active class in the selected school year.",
        )
    candidates = english.get(key, [])
    if len(candidates) == 1:
        return candidates[0], None
    if len(candidates) > 1:
        return None, _issue(
            "student_import_class_ambiguous",
            "class_name",
            "The English class label matches more than one class; use the full French stream name.",
        )
    return None, _issue(
        "student_import_class_not_found",
        "class_name",
        "No active class with this name exists in the selected school year.",
    )


def build_student_import_preview(
    db: Session,
    file_bytes: bytes,
    school_year: str,
    selected_rows: set[int] | None = None,
) -> dict:
    school_year = school_year.strip()
    if not school_year:
        raise StudentImportFileError("student_import_school_year_required", "School year is required.")
    parsed_rows = parse_student_workbook(file_bytes)
    if selected_rows is not None:
        parsed_rows = [row for row in parsed_rows if row["row_number"] in selected_rows]
    french_classes, english_classes = _class_maps(db, school_year)

    supplied_numbers = [row["student_number"] for row in parsed_rows if row["student_number"]]
    duplicate_numbers = {number for number, count in Counter(supplied_numbers).items() if count > 1}
    existing_numbers = set()
    if supplied_numbers:
        existing_numbers = set(
            db.scalars(select(Student.student_number).where(Student.student_number.in_(set(supplied_numbers)))).all()
        )

    names_by_email: dict[str, set[str]] = defaultdict(set)
    phones_by_email: dict[str, set[str]] = defaultdict(set)
    for row in parsed_rows:
        email = _canonical_email(row["parent_email"])
        if email:
            names_by_email[email].add(row["parent_name"].casefold())
            phones_by_email[email].add(row["parent_phone"])

    active_users: dict[str, User | None] = {}
    first_valid_row_for_email: dict[str, int] = {}
    preview_rows: list[dict] = []
    for row in parsed_rows:
        errors: list[dict] = []
        warnings: list[dict] = []
        for field in REQUIRED_COLUMNS:
            if not row[field]:
                errors.append(
                    _issue("student_import_required", field, f"{field} is required.")
                )

        email = _canonical_email(row["parent_email"])
        row["parent_email"] = email
        if email and not EMAIL_PATTERN.fullmatch(email):
            errors.append(
                _issue("student_import_invalid_email", "parent_email", "Parent email is invalid.")
            )
        if row["student_number"] and (
            row["student_number"] in existing_numbers or row["student_number"] in duplicate_numbers
        ):
            errors.append(
                _issue(
                    "student_import_duplicate_student_number",
                    "student_number",
                    "Student number is already used in the database or workbook.",
                )
            )
        if email and len(names_by_email[email]) > 1:
            errors.append(
                _issue(
                    "student_import_parent_name_conflict",
                    "parent_name",
                    "The same parent email has conflicting names in this workbook.",
                )
            )
        if email and len(phones_by_email[email]) > 1:
            warnings.append(
                _issue(
                    "student_import_parent_details_differ",
                    "parent_phone",
                    "The same parent email has different phone values; only the first new account value is kept.",
                )
            )

        school_class = None
        if row["class_name"]:
            school_class, class_error = _resolve_class(row["class_name"], french_classes, english_classes)
            if class_error:
                errors.append(class_error)

        parent_action = "create"
        active_user = None
        if email:
            if email not in active_users:
                active_users[email] = _active_user_by_canonical_email(db, email)
            active_user = active_users[email]
            if active_user is not None:
                parent = active_user.parent_profile
                if active_user.role != "parent" or parent is None or parent.deleted_at is not None:
                    errors.append(
                        _issue(
                            "student_import_email_role_conflict",
                            "parent_email",
                            "This email belongs to an active non-parent account.",
                        )
                    )
                else:
                    parent_action = "reuse"
                    if active_user.name.strip().casefold() != row["parent_name"].strip().casefold():
                        warnings.append(
                            _issue(
                                "student_import_parent_details_differ",
                                "parent_name",
                                "The existing parent account name will be kept.",
                            )
                        )
                    if row["parent_phone"] and (parent.phone or "").strip() != row["parent_phone"]:
                        warnings.append(
                            _issue(
                                "student_import_parent_details_differ",
                                "parent_phone",
                                "The existing parent account phone will be kept.",
                            )
                        )
            elif email in first_valid_row_for_email:
                parent_action = "reuse"

        is_valid = not errors
        if is_valid and email and active_user is None and email not in first_valid_row_for_email:
            first_valid_row_for_email[email] = row["row_number"]

        preview_rows.append(
            {
                **row,
                "class_id": school_class.id if school_class else None,
                "class_name": school_class.name_fr if school_class else row["class_name"],
                "school_level": school_class.school_level if school_class else None,
                "parent_action": parent_action,
                "is_valid": is_valid,
                "errors": errors,
                "warnings": warnings,
            }
        )

    valid_rows = sum(1 for row in preview_rows if row["is_valid"])
    return {
        "school_year": school_year,
        "total_rows": len(preview_rows),
        "valid_rows": valid_rows,
        "invalid_rows": len(preview_rows) - valid_rows,
        "rows": preview_rows,
    }


def generate_student_number(db: Session, year: int) -> str:
    prefix = f"STU-{year}-"
    numbers = db.scalars(
        select(Student.student_number).where(Student.student_number.like(f"{prefix}%"))
    ).all()
    suffixes = []
    for value in numbers:
        try:
            suffixes.append(int(value[len(prefix):]))
        except (TypeError, ValueError):
            continue
    number = max(suffixes, default=0) + 1
    while True:
        candidate = f"{prefix}{number:04d}"
        if db.scalar(select(Student.id).where(Student.student_number == candidate)) is None:
            return candidate
        number += 1


def _create_import_row(
    db: Session,
    preview_row: dict,
    current_user: User,
    parent_cache: dict[str, Parent],
) -> tuple[dict, tuple[Parent, str] | None, bool]:
    email = preview_row["parent_email"]
    existing_user = _active_user_by_canonical_email(db, email)
    parent = existing_user.parent_profile if existing_user and existing_user.role == "parent" else None
    parent_created = False
    notification: tuple[Parent, str] | None = None

    if email in parent_cache:
        parent = parent_cache[email]
    elif parent is None:
        temp_password = secrets.token_urlsafe(9)
        user = User(
            name=preview_row["parent_name"],
            email=email,
            password_hash=hash_password(temp_password),
            role="parent",
            must_change_password=True,
        )
        parent = Parent(user=user, phone=preview_row["parent_phone"] or None)
        db.add_all([user, parent])
        db.flush()
        create_audit_log(
            db=db,
            actor_user_id=current_user.id,
            action="parent_created",
            entity_type="parent",
            entity_id=parent.id,
            old_value=None,
            new_value={
                "user_id": parent.user_id,
                "name": user.name,
                "email": user.email,
                "phone": parent.phone,
                "source": "student_import",
            },
        )
        parent_created = True
        notification = (parent, temp_password)
    parent_cache[email] = parent

    student_number = preview_row["student_number"] or generate_student_number(
        db, datetime.now(timezone.utc).year
    )
    school_class = db.get(Class, preview_row["class_id"])
    if school_class is None or school_class.deleted_at is not None:
        raise ValueError("Class is no longer available")
    student = Student(
        first_name=preview_row["student_first_name"],
        last_name=preview_row["student_last_name"],
        school_level=school_class.school_level,
        student_number=student_number,
        educmaster_number=preview_row["educmaster_number"] or None,
        class_id=school_class.id,
    )
    db.add(student)
    db.flush()
    set_student_assignment(db, student, school_class)
    create_audit_log(
        db=db,
        actor_user_id=current_user.id,
        action="student_created",
        entity_type="student",
        entity_id=student.id,
        old_value=None,
        new_value={
            "first_name": student.first_name,
            "last_name": student.last_name,
            "school_level": student.school_level,
            "student_number": student.student_number,
            "educmaster_number": student.educmaster_number,
            "class_id": student.class_id,
            "source": "student_import",
        },
    )
    link = StudentParent(
        student_id=student.id,
        parent_id=parent.id,
        relationship=preview_row["relationship"] or None,
    )
    db.add(link)
    db.flush()
    create_audit_log(
        db=db,
        actor_user_id=current_user.id,
        action="parent_linked_to_student",
        entity_type="student_parent",
        entity_id=link.id,
        old_value=None,
        new_value={
            "student_id": student.id,
            "parent_id": parent.id,
            "relationship": link.relationship,
            "source": "student_import",
        },
    )
    return (
        {
            "row_number": preview_row["row_number"],
            "status": "created",
            "student_id": student.id,
            "student_number": student.student_number,
            "parent_id": parent.id,
            "parent_action": "create" if parent_created else "reuse",
        },
        notification,
        parent_created,
    )


def commit_student_import(
    db: Session,
    file_bytes: bytes,
    school_year: str,
    selected_rows: set[int],
    current_user: User,
) -> dict:
    parsed_rows = parse_student_workbook(file_bytes)
    available_rows = {row["row_number"] for row in parsed_rows}
    if not selected_rows:
        raise StudentImportFileError("student_import_no_rows_selected", "Select at least one row.")
    if not selected_rows.issubset(available_rows):
        raise StudentImportFileError(
            "student_import_invalid_selection",
            "Selected rows do not match the uploaded workbook.",
        )
    preview = build_student_import_preview(db, file_bytes, school_year, selected_rows=selected_rows)

    parent_cache: dict[str, Parent] = {}
    new_parent_emails: set[str] = set()
    reused_parent_emails: set[str] = set()
    notifications: dict[str, tuple[Parent, str]] = {}
    result_rows: list[dict] = []
    failures: list[dict] = []

    for preview_row in preview["rows"]:
        if preview_row["row_number"] not in selected_rows:
            continue
        if not preview_row["is_valid"]:
            failure = {
                "row_number": preview_row["row_number"],
                "errors": preview_row["errors"],
            }
            failures.append(failure)
            result_rows.append({"row_number": preview_row["row_number"], "status": "failed"})
            continue

        cached_emails_before = set(parent_cache)
        try:
            with db.begin_nested():
                result, notification, parent_created = _create_import_row(
                    db, preview_row, current_user, parent_cache
                )
            result_rows.append(result)
            email = preview_row["parent_email"]
            if parent_created:
                new_parent_emails.add(email)
                reused_parent_emails.discard(email)
                if notification is not None:
                    notifications[email] = notification
            elif email not in new_parent_emails:
                reused_parent_emails.add(email)
        except (IntegrityError, ValueError) as exc:
            for email in set(parent_cache) - cached_emails_before:
                parent_cache.pop(email, None)
            failure = {
                "row_number": preview_row["row_number"],
                "errors": [
                    _issue(
                        "student_import_commit_conflict",
                        None,
                        "The row conflicts with current data and could not be imported.",
                    )
                ],
            }
            failures.append(failure)
            result_rows.append({"row_number": preview_row["row_number"], "status": "failed"})

    db.commit()

    emails_sent = 0
    emails_failed = 0
    email_results: list[dict] = []
    for email, (parent, temp_password) in notifications.items():
        try:
            result = send_account_created_email(
                name=parent.user.name,
                to_email=email,
                temp_password=temp_password,
            )
        except Exception as exc:
            result = {"success": False, "error": str(exc)[:300] or exc.__class__.__name__}
        success = bool(result.get("success"))
        emails_sent += int(success)
        emails_failed += int(not success)
        email_results.append(
            {
                "parent_id": parent.id,
                "email": email,
                "success": success,
                "error": result.get("error"),
                "temp_password": None if success else temp_password,
            }
        )

    created = sum(1 for row in result_rows if row["status"] == "created")
    return {
        "status": "ok",
        "created": created,
        "failed": len(failures),
        "parents_reused": len(reused_parent_emails),
        "parents_created": len(new_parent_emails),
        "emails_sent": emails_sent,
        "emails_failed": emails_failed,
        "rows": result_rows,
        "failures": failures,
        "email_results": email_results,
    }
