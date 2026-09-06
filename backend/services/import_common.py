"""Shared XLSX parsing primitives for explicit preview-then-commit imports.

Domain validation stays in each importer. This module only enforces the file
envelope so teacher, parent-link, and future imports agree on header safety,
row limits, normalization, and actionable error codes.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from typing import Any

from openpyxl import load_workbook


MAX_IMPORT_ROWS = 500
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class ImportFileError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def canonical_email(value: Any) -> str:
    return text(value).lower()


def issue(code: str, field: str | None, message: str) -> dict:
    return {"code": code, "field": field, "message": message}


def parse_workbook(
    file_bytes: bytes,
    *,
    required_columns: tuple[str, ...],
    optional_columns: tuple[str, ...],
    code_prefix: str,
    row_label: str,
) -> list[dict]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(
            f"{code_prefix}_invalid_workbook",
            "The uploaded file is not a valid .xlsx workbook.",
        ) from exc

    try:
        sheet = workbook.active
        raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if raw_headers is None:
            raise ImportFileError(f"{code_prefix}_empty_workbook", "The workbook is empty.")

        headers = [text(value) for value in raw_headers]
        nonempty_headers = [header for header in headers if header]
        duplicates = sorted(header for header, count in Counter(nonempty_headers).items() if count > 1)
        if duplicates:
            raise ImportFileError(
                f"{code_prefix}_duplicate_columns",
                f"Duplicate columns: {', '.join(duplicates)}",
            )

        missing = [column for column in required_columns if column not in nonempty_headers]
        if missing:
            raise ImportFileError(
                f"{code_prefix}_missing_columns",
                f"Missing required columns: {', '.join(missing)}",
            )

        allowed_columns = set(required_columns + optional_columns)
        unknown = sorted(header for header in nonempty_headers if header not in allowed_columns)
        if unknown:
            raise ImportFileError(
                f"{code_prefix}_unknown_columns",
                f"Unsupported columns: {', '.join(unknown)}",
            )

        rows: list[dict] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if any(
                not headers[index] and text(value)
                for index, value in enumerate(values)
                if index < len(headers)
            ):
                raise ImportFileError(
                    f"{code_prefix}_unknown_columns",
                    "Every populated column must have a supported heading.",
                )
            mapped = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            if not any(text(value) for value in mapped.values()):
                continue
            rows.append(
                {
                    "row_number": row_number,
                    **{column: text(mapped.get(column)) for column in allowed_columns},
                }
            )
            if len(rows) > MAX_IMPORT_ROWS:
                raise ImportFileError(
                    f"{code_prefix}_too_many_rows",
                    f"A workbook may contain at most {MAX_IMPORT_ROWS} data rows.",
                )

        if not rows:
            raise ImportFileError(
                f"{code_prefix}_empty_workbook",
                f"The workbook has no {row_label} rows.",
            )
        return rows
    finally:
        workbook.close()
